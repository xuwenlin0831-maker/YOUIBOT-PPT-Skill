#!/usr/bin/env python3
"""Validate slide count, editability signals, and unresolved template text."""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


PLACEHOLDERS = (
    "点击输入",
    "您的内容",
    "输入内容",
    "添加内容",
    "添加标题",
    "请输入您的小标题",
    "副标题",
    "内容范围",
)


def iter_shapes(shapes):
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from iter_shapes(shape.shapes)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("pptx", type=Path)
    parser.add_argument("--expected-slides", type=int)
    args = parser.parse_args()

    deck = Presentation(args.pptx.resolve())
    issues = []
    if args.expected_slides is not None and len(deck.slides) != args.expected_slides:
        issues.append(
            f"expected {args.expected_slides} slides, found {len(deck.slides)}"
        )

    details = []
    for number, slide in enumerate(deck.slides, start=1):
        counts = {"text": 0, "pictures": 0, "tables": 0, "charts": 0, "shapes": 0}
        unresolved = []
        for shape in iter_shapes(slide.shapes):
            counts["shapes"] += 1
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                counts["pictures"] += 1
            if getattr(shape, "has_table", False):
                counts["tables"] += 1
                for row in shape.table.rows:
                    for cell in row.cells:
                        for token in PLACEHOLDERS:
                            if token in cell.text:
                                unresolved.append(token)
            if getattr(shape, "has_chart", False):
                counts["charts"] += 1
            if getattr(shape, "has_text_frame", False) and shape.text.strip():
                counts["text"] += 1
                for token in PLACEHOLDERS:
                    if token in shape.text:
                        unresolved.append(token)
        if unresolved:
            issues.append(
                f"slide {number} contains unresolved template text: "
                + ", ".join(sorted(set(unresolved)))
            )
        details.append({"slide": number, **counts})

    report = {
        "file": str(args.pptx.resolve()),
        "slides": len(deck.slides),
        "editable_object_counts": details,
        "issues": issues,
        "ok": not issues,
    }
    print(json.dumps(report, ensure_ascii=False, indent=2))
    raise SystemExit(1 if issues else 0)


if __name__ == "__main__":
    main()
